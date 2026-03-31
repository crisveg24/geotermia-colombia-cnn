"""
Genera un archivo Word (.docx) con las 15 ecuaciones de la tesis
ya renderizadas en el editor de ecuaciones nativo de Word.

Uso:
    python docs/generar_ecuaciones_word.py

Se genera: docs/ecuaciones_tesis.docx
"""

import latex2mathml.converter
from lxml import etree
from docx import Document
from docx.shared import Pt
from pathlib import Path
import copy
import re

# Ruta al XSLT de Microsoft Office (convierte MathML → OMML)
XSLT_PATH = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"

# Directorio de salida
OUT_DIR = Path(__file__).parent
OUT_FILE = OUT_DIR / "ecuaciones_tesis.docx"

# Las 15 ecuaciones del documento TESIS.md
ECUACIONES = [
    {
        "num": 1,
        "nombre": "Convolución (Conv2D)",
        "seccion": "§5.3.2",
        "latex": r"\text{Output}(i,j) = \sum_{m,n} \text{Input}(i+m, j+n) \times \text{Kernel}(m,n) + b",
    },
    {
        "num": 2,
        "nombre": "Activación ReLU",
        "seccion": "§5.3.2",
        "latex": r"f(x) = \max(0, x)",
    },
    {
        "num": 3,
        "nombre": "Parámetro (peso y sesgo)",
        "seccion": "§5.3.3",
        "latex": r"\text{salida} = \text{entrada} \times w + b",
    },
    {
        "num": 4,
        "nombre": "Función sigmoid",
        "seccion": "§5.3.4",
        "latex": r"\sigma(x) = \frac{1}{1 + e^{-x}}",
    },
    {
        "num": 5,
        "nombre": "Binary Cross-Entropy",
        "seccion": "§5.3.5",
        "latex": r"\mathcal{L} = -\left[ y \cdot \log(\hat{y}) + (1 - y) \cdot \log(1 - \hat{y}) \right]",
    },
    {
        "num": 6,
        "nombre": "Conexión residual",
        "seccion": "§5.4.2",
        "latex": r"y = F(x, \{W_i\}) + x",
    },
    {
        "num": 7,
        "nombre": "MixUp",
        "seccion": "§5.5",
        "latex": r"\tilde{x} = \lambda x_i + (1 - \lambda) x_j, \quad \tilde{y} = \lambda y_i + (1 - \lambda) y_j",
    },
    {
        "num": 8,
        "nombre": "Normalización z-score global",
        "seccion": "§6.3",
        "latex": r"x_{\text{norm}} = \frac{x - \mu_{\text{banda}}^{\text{global}}}{\sigma_{\text{banda}}^{\text{global}}}",
    },
    {
        "num": 9,
        "nombre": "Hipótesis nula",
        "seccion": "§4.2",
        "latex": r"H_0: \text{Accuracy} \leq 0{,}50",
    },
    {
        "num": 10,
        "nombre": "Accuracy",
        "seccion": "§6.6",
        "latex": r"\text{Accuracy} = \frac{TP + TN}{TP + TN + FP + FN}",
    },
    {
        "num": 11,
        "nombre": "Precision",
        "seccion": "§6.6",
        "latex": r"\text{Precision} = \frac{TP}{TP + FP}",
    },
    {
        "num": 12,
        "nombre": "Recall",
        "seccion": "§6.6",
        "latex": r"\text{Recall} = \frac{TP}{TP + FN}",
    },
    {
        "num": 13,
        "nombre": "F1-Score",
        "seccion": "§6.6",
        "latex": r"F_1 = 2 \times \frac{\text{Precision} \times \text{Recall}}{\text{Precision} + \text{Recall}}",
    },
    {
        "num": 14,
        "nombre": "MCC (Matthews Correlation Coefficient)",
        "seccion": "§6.6",
        "latex": r"\text{MCC} = \frac{TP \times TN - FP \times FN}{\sqrt{(TP+FP)(TP+FN)(TN+FP)(TN+FN)}}",
    },
    {
        "num": 15,
        "nombre": "Intervalo de confianza Bootstrap",
        "seccion": "§6.7",
        "latex": r"\text{IC}_{1-\alpha} = \left[\hat{\theta}^{*}_{(\alpha/2)},\; \hat{\theta}^{*}_{(1-\alpha/2)}\right]",
    },
]


def latex_to_omml(latex_str: str):
    """Convierte LaTeX → MathML → OMML (Office Math Markup Language)."""
    # latex2mathml convierte LaTeX a MathML
    mathml_str = latex2mathml.converter.convert(latex_str)

    # Parsear el MathML
    mathml_tree = etree.fromstring(mathml_str.encode("utf-8"))

    # Cargar la hoja XSLT de Microsoft Office
    xslt_tree = etree.parse(XSLT_PATH)
    transform = etree.XSLT(xslt_tree)

    # Transformar MathML → OMML
    omml_tree = transform(mathml_tree)

    return omml_tree.getroot()


def main():
    # Verificar que existe el XSLT
    if not Path(XSLT_PATH).exists():
        print(f"ERROR: No se encontró {XSLT_PATH}")
        print("Necesitas tener Microsoft Office instalado.")
        return

    doc = Document()

    # Título
    title = doc.add_heading("Ecuaciones de la Tesis", level=1)

    doc.add_paragraph(
        "Este documento contiene las 15 ecuaciones del proyecto de grado "
        "ya formateadas en el editor de ecuaciones de Word. "
        "Para usarlas: selecciona la ecuación, cópiala (Ctrl+C) y pégala "
        "en el documento de Word Online (Ctrl+V)."
    )
    doc.add_paragraph("")

    # Generar cada ecuación
    for eq in ECUACIONES:
        # Etiqueta descriptiva
        label = f"Ecuación ({eq['num']}) — {eq['nombre']} [{eq['seccion']}]"
        p_label = doc.add_paragraph()
        run = p_label.add_run(label)
        run.bold = True
        run.font.size = Pt(11)

        # Párrafo con la ecuación OMML
        p_eq = doc.add_paragraph()
        p_eq.alignment = 1  # Center

        try:
            omml_element = latex_to_omml(eq["latex"])
            # Insertar el OMML en el párrafo
            p_eq._element.append(omml_element)
            print(f"  ✓ Ecuación ({eq['num']}) — {eq['nombre']}")
        except Exception as e:
            # Si falla, poner el LaTeX como texto plano
            p_eq.add_run(f"[LaTeX] {eq['latex']}")
            print(f"  ✗ Ecuación ({eq['num']}) falló: {e}")

        # Número de ecuación a la derecha
        p_num = doc.add_paragraph(f"({eq['num']})")
        p_num.alignment = 2  # Right

        # Separador
        doc.add_paragraph("")

    # Guardar
    doc.save(str(OUT_FILE))
    print(f"\n✓ Archivo generado: {OUT_FILE}")
    print("  Ábrelo en Word desktop, selecciona cada ecuación y cópiala.")


if __name__ == "__main__":
    main()
